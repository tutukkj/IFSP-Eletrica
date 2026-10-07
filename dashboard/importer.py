import math
import unicodedata
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.dialects.sqlite import insert as sqlite_insert
from .database import metadata, readings
from .metrics import METRICS, is_valid


def normalize(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value or '').strip().lower()) if not unicodedata.combining(c))


def parse_number(value):
    if value is None or str(value).strip() == '':
        return None
    result = float(str(value).strip().replace(',', '.'))
    if not math.isfinite(result):
        raise ValueError('Número não finito')
    return result


def parse_timestamp(day, clock):
    if isinstance(day, datetime):
        day = day.date()
    elif not isinstance(day, date):
        day = datetime.strptime(str(day).strip(), '%d/%m/%Y').date()
    if isinstance(clock, datetime):
        clock = clock.time()
    elif not isinstance(clock, time):
        clock = time.fromisoformat(str(clock).strip())
    if clock.tzinfo:
        raise ValueError('O horário deve estar sem fuso, como na planilha original.')
    return datetime.combine(day, clock)


def import_excel(engine, filename, station='principal', sheet=None, progress=None):
    station = station.strip()
    if not station or len(station) > 120:
        raise ValueError('Estação deve ter entre 1 e 120 caracteres.')
    path = Path(filename)
    if len(path.name) > 255:
        raise ValueError('Nome do arquivo muito longo.')
    workbook = load_workbook(path, read_only=True, data_only=True)
    stats = {'read': 0, 'inserted': 0, 'duplicates': 0, 'quality_rows': 0}
    try:
        ws = workbook[sheet] if sheet else workbook.worksheets[0]
        # Este XLSX declara A1 como dimensão embora tenha mais de 65 mil linhas.
        ws.reset_dimensions()
        iterator = ws.iter_rows(values_only=True)
        header = next(iterator)
        positions = {normalize(name): i for i, name in enumerate(header)}
        required = ['DATE', 'TIME', *(v[0] for v in METRICS.values())]
        missing = [name for name in required if normalize(name) not in positions]
        if missing:
            raise ValueError('Colunas ausentes: ' + ', '.join(missing))
        metadata.create_all(engine)
        insert = pg_insert if engine.dialect.name == 'postgresql' else sqlite_insert
        statement = insert(readings).on_conflict_do_nothing(index_elements=['station', 'timestamp']).returning(readings.c.station)
        with engine.begin() as connection:
            batch = []
            def flush():
                if batch:
                    result = connection.execute(statement, batch)
                    stats['inserted'] += len(result.fetchall())
                    batch.clear()
                    if progress:
                        progress(stats['read'])
            for line, values in enumerate(iterator, 2):
                if not any(v is not None for v in values):
                    continue
                def get(name):
                    index = positions[normalize(name)]
                    return values[index] if index < len(values) else None
                try:
                    row = {'station': station, 'timestamp': parse_timestamp(get('DATE'), get('TIME')), 'source_file': path.name}
                    row.update({key: parse_number(get(spec[0])) for key, spec in METRICS.items()})
                    row['has_quality_issue'] = any(not is_valid(key, row[key], row) for key in METRICS)
                except (ValueError, TypeError, OverflowError) as exc:
                    raise ValueError(f'Linha {line}: {exc}') from exc
                stats['read'] += 1
                stats['quality_rows'] += int(row['has_quality_issue'])
                batch.append(row)
                if len(batch) >= 500:
                    flush()
            flush()
        stats['duplicates'] = stats['read'] - stats['inserted']
        return stats
    finally:
        workbook.close()
